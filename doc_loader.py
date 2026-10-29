"""
Document Loader (Word .docx & Text .txt / .md)
Extracts headings, paragraphs, and tables from Word documents and plain text files.
"""

import os
import io
from typing import List, Dict, Any, Union
import docx


class DocChunk:
    """Represents a text chunk from a Word or Text document."""
    def __init__(self, doc_id: str, section: str, text: str, format_type: str = "text"):
        self.doc_id = doc_id
        self.section = section
        self.text = text
        self.format_type = format_type

    @property
    def citation(self) -> str:
        sec_str = f" | Section: '{self.section}'" if self.section else ""
        return f"[Source: {self.doc_id}{sec_str}]"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "doc_id": self.doc_id,
            "section": self.section,
            "text": self.text,
            "format_type": self.format_type,
            "citation": self.citation
        }


class DocumentLoader:
    """Loads DOCX, TXT, and Markdown files."""

    @staticmethod
    def load(file_input: Union[str, bytes, io.BytesIO], filename: str, chunk_size: int = 800, overlap: int = 100) -> List[DocChunk]:
        """Parse document and return chunks."""
        ext = os.path.splitext(filename)[1].lower()

        if ext == ".docx":
            return DocumentLoader._load_docx(file_input, filename, chunk_size, overlap)
        else:
            return DocumentLoader._load_text(file_input, filename, chunk_size, overlap)

    @staticmethod
    def _load_docx(file_input: Any, filename: str, chunk_size: int, overlap: int) -> List[DocChunk]:
        """Parse docx file."""
        if isinstance(file_input, (str, os.PathLike)):
            doc = docx.Document(file_input)
        elif isinstance(file_input, bytes):
            doc = docx.Document(io.BytesIO(file_input))
        elif hasattr(file_input, "read"):
            data = file_input.read()
            doc = docx.Document(io.BytesIO(data if isinstance(data, bytes) else data.encode('utf-8')))
        else:
            raise ValueError("Unsupported DOCX input")

        blocks = []
        current_heading = "Introduction"

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                current_heading = text
            else:
                blocks.append((current_heading, text))

        # Also extract tables from docx
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_vals = [cell.text.strip() for cell in row.cells]
                table_lines.append(" | ".join(row_vals))
            if table_lines:
                blocks.append((current_heading, "\n".join(table_lines)))

        return DocumentLoader._assemble_chunks(blocks, filename, "docx", chunk_size, overlap)

    @staticmethod
    def _load_text(file_input: Any, filename: str, chunk_size: int, overlap: int) -> List[DocChunk]:
        """Parse plain text or markdown."""
        if isinstance(file_input, (str, os.PathLike)) and os.path.exists(str(file_input)):
            with open(file_input, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif isinstance(file_input, str):
            content = file_input
        elif isinstance(file_input, bytes):
            content = file_input.decode("utf-8", errors="replace")
        elif hasattr(file_input, "read"):
            data = file_input.read()
            content = data.decode("utf-8", errors="replace") if isinstance(data, bytes) else str(data)
        else:
            content = str(file_input)

        paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
        blocks = [("Text Section", p) for p in paragraphs]
        return DocumentLoader._assemble_chunks(blocks, filename, "text", chunk_size, overlap)

    @staticmethod
    def _assemble_chunks(blocks: List[tuple], filename: str, fmt: str, chunk_size: int, overlap: int) -> List[DocChunk]:
        chunks = []
        curr_text = []
        curr_len = 0
        curr_sec = "Overview"

        for sec, text in blocks:
            if curr_len + len(text) > chunk_size and curr_text:
                chunks.append(DocChunk(
                    doc_id=filename,
                    section=curr_sec,
                    text="\n\n".join(curr_text),
                    format_type=fmt
                ))
                curr_text = [text]
                curr_len = len(text)
                curr_sec = sec
            else:
                curr_text.append(text)
                curr_len += len(text)
                curr_sec = sec

        if curr_text:
            chunks.append(DocChunk(
                doc_id=filename,
                section=curr_sec,
                text="\n\n".join(curr_text),
                format_type=fmt
            ))

        return chunks
